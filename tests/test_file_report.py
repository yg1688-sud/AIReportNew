"""Tests for file-based report generators (file_excel.py and file_pdf.py)."""

from pathlib import Path

import pytest


class TestFileExcelReport:
    """Unit tests for generate_file_excel_report()."""

    def test_generates_valid_file(self, sample_file_df, temp_dir):
        """Generates a valid .xlsx with expected structure."""
        from src.report.file_excel import generate_file_excel_report

        path = generate_file_excel_report(
            data=sample_file_df,
            value_col="销售金额",
            output_dir=str(temp_dir),
            file_label="test_file",
        )

        assert Path(path).exists()
        assert Path(path).suffix == ".xlsx"

        from openpyxl import load_workbook
        wb = load_workbook(path)
        ws = wb.active
        assert ws.max_row >= 3  # title + summary + header + data + total

    def test_contains_all_original_columns(self, sample_file_df, temp_dir):
        """Excel headers include all original column names."""
        from src.report.file_excel import generate_file_excel_report

        path = generate_file_excel_report(
            data=sample_file_df,
            value_col="销售金额",
            output_dir=str(temp_dir),
            file_label="test_file",
        )

        from openpyxl import load_workbook
        wb = load_workbook(path)
        ws = wb.active
        headers = [ws.cell(7, c).value for c in range(1, ws.max_column + 1)]
        for col in sample_file_df.columns:
            assert col in headers

    def test_contains_percentage_column(self, sample_file_df, temp_dir):
        """A '占比' column is appended."""
        from src.report.file_excel import generate_file_excel_report

        path = generate_file_excel_report(
            data=sample_file_df,
            value_col="销售金额",
            output_dir=str(temp_dir),
            file_label="test_file",
        )

        from openpyxl import load_workbook
        wb = load_workbook(path)
        ws = wb.active
        headers = [ws.cell(7, c).value for c in range(1, ws.max_column + 1)]
        assert "占比" in headers

    def test_money_columns_formatted(self, sample_file_df, temp_dir):
        """Money columns use number format."""
        from src.report.file_excel import generate_file_excel_report

        path = generate_file_excel_report(
            data=sample_file_df,
            value_col="销售金额",
            output_dir=str(temp_dir),
            file_label="test_file",
        )

        from openpyxl import load_workbook
        wb = load_workbook(path)
        ws = wb.active
        # Find the "销售金额" column index
        headers = [ws.cell(7, c).value for c in range(1, ws.max_column + 1)]
        idx = headers.index("销售金额") + 1
        # First data row (row 8 after header at row 7) should have number_format set
        cell = ws.cell(8, idx)
        assert cell.number_format == "#,##0.00"

    def test_total_row_present(self, sample_file_df, temp_dir):
        """Last row is a total row with '合计' label."""
        from src.report.file_excel import generate_file_excel_report

        path = generate_file_excel_report(
            data=sample_file_df,
            value_col="销售金额",
            output_dir=str(temp_dir),
            file_label="test_file",
        )

        from openpyxl import load_workbook
        wb = load_workbook(path)
        ws = wb.active
        last_row = ws.max_row
        assert ws.cell(last_row, 1).value == "合计"

    def test_empty_dataframe_handled(self, temp_dir):
        """Empty DataFrame generates minimal valid file."""
        import pandas as pd
        from src.report.file_excel import generate_file_excel_report

        df = pd.DataFrame({"金额": []})
        path = generate_file_excel_report(
            data=df,
            value_col="金额",
            output_dir=str(temp_dir),
            file_label="empty_test",
        )
        assert Path(path).exists()

    def test_percentage_column_name_conflict(self, sample_file_df, temp_dir):
        """If '占比' already exists, computed column is '占比(分析)'."""
        from src.report.file_excel import generate_file_excel_report

        df = sample_file_df.copy()
        df["占比"] = "10%"

        path = generate_file_excel_report(
            data=df,
            value_col="销售金额",
            output_dir=str(temp_dir),
            file_label="test_conflict",
        )

        from openpyxl import load_workbook
        wb = load_workbook(path)
        ws = wb.active
        headers = [ws.cell(7, c).value for c in range(1, ws.max_column + 1)]
        assert "占比(分析)" in headers
        assert "占比(原始)" in headers


class TestFilePDFReport:
    """Unit tests for generate_file_pdf_report()."""

    def test_generates_valid_file(self, sample_file_df, temp_dir):
        """Generates a valid .pdf file."""
        from src.report.file_pdf import generate_file_pdf_report

        path = generate_file_pdf_report(
            data=sample_file_df,
            value_col="销售金额",
            output_dir=str(temp_dir),
            file_label="test_file",
        )

        assert Path(path).exists()
        assert Path(path).suffix == ".pdf"
        # PDF magic bytes
        with open(path, "rb") as f:
            header = f.read(5)
        assert header == b"%PDF-"

    def test_handles_empty_data(self, temp_dir):
        """Empty DataFrame generates valid PDF."""
        import pandas as pd
        from src.report.file_pdf import generate_file_pdf_report

        df = pd.DataFrame({"金额": []})
        path = generate_file_pdf_report(
            data=df,
            value_col="金额",
            output_dir=str(temp_dir),
            file_label="empty_pdf",
        )

        assert Path(path).exists()
        with open(path, "rb") as f:
            header = f.read(5)
        assert header == b"%PDF-"

    def test_handles_many_columns(self, temp_dir):
        """DataFrame with many columns generates valid PDF."""
        import pandas as pd
        from src.report.file_pdf import generate_file_pdf_report

        df = pd.DataFrame({
            f"列{i}": [i * 10, i * 20, i * 30] for i in range(1, 9)
        })
        df["金额"] = [100, 200, 300]

        path = generate_file_pdf_report(
            data=df,
            value_col="金额",
            output_dir=str(temp_dir),
            file_label="wide_test",
        )

        assert Path(path).exists()
        with open(path, "rb") as f:
            header = f.read(5)
        assert header == b"%PDF-"
