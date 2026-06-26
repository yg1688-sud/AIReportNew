"""Generate professional project presentation PPT for AIExport system.

Design: dark-slate theme with blue accents, 16:9 format, 12 slides.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ── Color Palette (dark professional theme) ──
BG_DARK      = RGBColor(0x1A, 0x1F, 0x2E)   # Deep navy background
BG_SECTION   = RGBColor(0x22, 0x28, 0x3A)   # Slightly lighter section bg
ACCENT_BLUE  = RGBColor(0x3B, 0x82, 0xF6)   # Primary blue
ACCENT_TEAL  = RGBColor(0x14, 0xB8, 0xA6)   # Teal
ACCENT_ORANGE= RGBColor(0xF5, 0x9E, 0x0B)   # Amber/orange
ACCENT_GREEN = RGBColor(0x10, 0xB9, 0x81)   # Emerald
ACCENT_PURPLE= RGBColor(0x8B, 0x5C, 0xF6)   # Purple
WHITE        = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY   = RGBColor(0x94, 0xA3, 0xB8)   # Slate-400
DARK_CARD    = RGBColor(0x1E, 0x29, 0x3B)   # Card background
BORDER_SUBTLE= RGBColor(0x33, 0x40, 0x55)   # Subtle border

OUTPUT_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "output")
PPT_PATH   = os.path.join(OUTPUT_DIR, "AIExport_项目讲解.pptx")


# ── Helpers ──

def rect(slide, l, t, w, h, color, radius=None):
    """Add a rounded or sharp rectangle."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        l, t, w, h
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    if radius:
        shape.adjustments[0] = radius
    return shape

def textbox(slide, l, t, w, h, text, size=18, color=WHITE, bold=False,
            align=PP_ALIGN.LEFT, font="Microsoft YaHei", anchor=MSO_ANCHOR.TOP):
    """Add a text box."""
    tb = slide.shapes.add_textbox(l, t, w, h)
    tb.word_wrap = True
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font
    p.alignment = align
    return tb

def bullets(slide, l, t, w, h, items, size=13, color=LIGHT_GRAY, spacing=1.6,
            font_name="Microsoft YaHei"):
    """Add bullet-point text block."""
    tb = slide.shapes.add_textbox(l, t, w, h)
    tb.word_wrap = True
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.name = font_name
        p.space_after = Pt(size * (spacing - 1) + 4)
    return tb

def header_bar(slide, title, subtitle=""):
    """Dark header bar at top."""
    rect(slide, Inches(0), Inches(0), Inches(13.33), Inches(1.3), BG_SECTION)
    # accent line
    rect(slide, Inches(0), Inches(1.3), Inches(13.33), Inches(0.04), ACCENT_BLUE)
    textbox(slide, Inches(0.8), Inches(0.2), Inches(11), Inches(0.7),
            title, size=32, color=WHITE, bold=True)
    if subtitle:
        textbox(slide, Inches(0.8), Inches(0.75), Inches(11), Inches(0.4),
                subtitle, size=14, color=LIGHT_GRAY)

def slide_num(slide, n, total=12):
    textbox(slide, Inches(11.5), Inches(7.0), Inches(1.5), Inches(0.3),
            f"{n} / {total}", size=9, color=LIGHT_GRAY, align=PP_ALIGN.RIGHT)

def card(slide, x, y, w, h, title, items, title_color=ACCENT_BLUE, icon=""):
    """A content card with title and bullet items."""
    rect(slide, x, y, w, h, DARK_CARD, radius=0.04)
    # top accent line
    rect(slide, x, y, w, Inches(0.04), title_color)
    header_text = f"{icon} {title}" if icon else title
    textbox(slide, x + Inches(0.25), y + Inches(0.2), w - Inches(0.5), Inches(0.45),
            header_text, size=16, color=title_color, bold=True)
    bullets(slide, x + Inches(0.25), y + Inches(0.7), w - Inches(0.5), h - Inches(0.9),
            items, size=11, color=LIGHT_GRAY, spacing=1.65)

def tag(slide, x, y, text, color=ACCENT_BLUE):
    """Small tag/badge."""
    s = rect(slide, x, y, Inches(1.5), Inches(0.32), color, radius=0.15)
    textbox(slide, x, y + Inches(0.02), Inches(1.5), Inches(0.28),
            text, size=9, color=WHITE, bold=True, align=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════
# SLIDE 1 — Cover
# ═══════════════════════════════════════════════
def slide_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # Full dark background
    rect(slide, Inches(0), Inches(0), Inches(13.33), Inches(7.5), BG_DARK)
    # Decorative gradient-like blocks
    rect(slide, Inches(0), Inches(0), Inches(13.33), Inches(0.08), ACCENT_BLUE)
    rect(slide, Inches(0), Inches(6.0), Inches(13.33), Inches(0.06), ACCENT_BLUE)
    # Right-side decorative element
    rect(slide, Inches(9.5), Inches(1.0), Inches(3.83), Inches(5.5), BG_SECTION)
    rect(slide, Inches(9.5), Inches(1.0), Inches(0.06), Inches(5.5), ACCENT_BLUE)

    # Main title
    textbox(slide, Inches(1.0), Inches(1.8), Inches(7.5), Inches(1.2),
            "AIExport", size=60, color=WHITE, bold=True)
    textbox(slide, Inches(1.0), Inches(3.0), Inches(7.5), Inches(0.7),
            "AI 智能导出与分析系统", size=30, color=ACCENT_BLUE)
    textbox(slide, Inches(1.0), Inches(3.7), Inches(7.5), Inches(0.5),
            "配置驱动 · 自动模版 · 多格式报告 · 可视化图表", size=15, color=LIGHT_GRAY)

    # Right side — key numbers
    metrics = [
        ("< 30s", "全流程耗时"),
        ("0 误差", "计算精度"),
        ("3 格式", "报告输出"),
        ("10+", "模版管理"),
    ]
    for i, (val, lab) in enumerate(metrics):
        y = Inches(1.6) + i * Inches(1.1)
        textbox(slide, Inches(10.0), y, Inches(2.8), Inches(0.5),
                val, size=36, color=ACCENT_BLUE, bold=True, align=PP_ALIGN.CENTER)
        textbox(slide, Inches(10.0), y + Inches(0.5), Inches(2.8), Inches(0.3),
                lab, size=12, color=LIGHT_GRAY, align=PP_ALIGN.CENTER)

    # Bottom info
    textbox(slide, Inches(1.0), Inches(6.3), Inches(7.0), Inches(0.35),
            "版本 v2.0  |  纯后端 CLI 工具  |  Python 3.11+  |  SQL Server + YAML 驱动",
            size=12, color=LIGHT_GRAY)
    slide_num(slide, 1)


# ═══════════════════════════════════════════════
# SLIDE 2 — Background & Pain Points
# ═══════════════════════════════════════════════
def slide_background(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, Inches(0), Inches(0), Inches(13.33), Inches(7.5), BG_DARK)
    header_bar(slide, "项目背景与痛点", "Why AIExport — 从手工到自动化的飞跃")

    # Left — Pain points
    textbox(slide, Inches(0.8), Inches(1.7), Inches(5.5), Inches(0.45),
            "🔴  传统手工流程的 6 大痛点", size=20, color=ACCENT_ORANGE, bold=True)
    pains = [
        "① 手工导出 — 登录 SSMS，执行 SQL，导出 CSV，重复繁琐",
        "② 手工匹配 — 数十个员工逐一核对，门店/片区/部门手工归类",
        "③ 手工计算 — Excel 公式易出错，跨表引用断裂，检查费时",
        "④ 手工制报告 — 反复调整格式、加粗、对齐，每次都要重做",
        "⑤ 无可视化 — 缺乏图表辅助决策，管理层看不到直观数据",
        "⑥ 耗时长 — 每次半月报/月报需 2~4 小时重复劳动",
    ]
    bullets(slide, Inches(0.8), Inches(2.3), Inches(5.5), Inches(3.5),
            pains, size=12, color=LIGHT_GRAY, spacing=1.7)

    # Right — Solution
    textbox(slide, Inches(7.0), Inches(1.7), Inches(5.5), Inches(0.45),
            "🟢  AIExport 一键解决方案", size=20, color=ACCENT_GREEN, bold=True)
    solutions = [
        "① 配置化查询 — YAML 定义 SQL，{{变量}} 参数化，自动执行",
        "② 智能模版 — 从查询结果自动生成分析模版，零手工配置",
        "③ 自动计算 — 分组汇总 + 占比计算，与 SQL 总额 0 偏差",
        "④ 一键报告 — Markdown + Excel + PDF 三格式一步到位",
        "⑤ 可视化图表 — 柱状图/饼图/漏斗图自动生成，中文完美",
        "⑥ 定时调度 — --auto-date 每月 1 号/16 号自动导出分析",
    ]
    bullets(slide, Inches(7.0), Inches(2.3), Inches(5.5), Inches(3.5),
            solutions, size=12, color=LIGHT_GRAY, spacing=1.7)

    # Bottom highlight bar
    rect(slide, Inches(0.6), Inches(6.4), Inches(12.1), Inches(0.7), DARK_CARD, radius=0.04)
    textbox(slide, Inches(0.8), Inches(6.5), Inches(11.7), Inches(0.5),
            "💡 核心价值：将 2~4 小时的手工流程压缩至 30 秒内全自动完成，零人工介入，100% 计算精确",
            size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    slide_num(slide, 2)


# ═══════════════════════════════════════════════
# SLIDE 3 — System Architecture Overview
# ═══════════════════════════════════════════════
def slide_architecture(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, Inches(0), Inches(0), Inches(13.33), Inches(7.5), BG_DARK)
    header_bar(slide, "系统架构总览", "Export → Analyze → Report → Visualize — 全流程自动化管道")

    # Pipeline flow
    stages = [
        ("📋 配置加载", "YAML 配置解析\n连接参数 + SQL\n多查询组管理", ACCENT_BLUE),
        ("📤 数据导出", "pymssql 连接\n参数化查询执行\n.xlsx 输出", ACCENT_TEAL),
        ("🧠 智能分析", "自动模版生成\n员工匹配 & 汇总\n占比精确计算", ACCENT_PURPLE),
        ("📊 报告生成", "Excel + PDF\nMarkdown\n图表 PNG", ACCENT_ORANGE),
    ]

    bw = Inches(2.6)
    bh = Inches(2.2)
    gap = Inches(0.45)
    sx = Inches(0.9)
    sy = Inches(1.8)

    for i, (title, desc, color) in enumerate(stages):
        x = sx + i * (bw + gap)
        rect(slide, x, sy, bw, bh, DARK_CARD, radius=0.06)
        rect(slide, x, sy, bw, Inches(0.05), color)
        textbox(slide, x + Inches(0.05), sy + Inches(0.15), bw - Inches(0.1), Inches(0.15),
                f"STEP {i+1}", size=9, color=color, align=PP_ALIGN.CENTER)
        textbox(slide, x + Inches(0.15), sy + Inches(0.45), bw - Inches(0.3), Inches(0.5),
                title, size=16, color=color, bold=True, align=PP_ALIGN.CENTER)
        textbox(slide, x + Inches(0.15), sy + Inches(1.1), bw - Inches(0.3), Inches(1.0),
                desc, size=11, color=LIGHT_GRAY, align=PP_ALIGN.CENTER)
        # arrow
        if i < len(stages) - 1:
            textbox(slide, x + bw, sy + Inches(0.7), gap, Inches(0.5),
                    "▸", size=24, color=ACCENT_BLUE, align=PP_ALIGN.CENTER)

    # Two working modes
    rect(slide, Inches(0.7), Inches(4.5), Inches(5.8), Inches(1.5), DARK_CARD, radius=0.06)
    textbox(slide, Inches(1.0), Inches(4.65), Inches(3.0), Inches(0.35),
            "🔹 模式一：多查询管道", size=15, color=ACCENT_BLUE, bold=True)
    mode1 = [
        "• SQL Server → 自动模版 → 分组分析 → 报告/图表",
        "• 适用：固定周期报表（半月报/月报），定时调度",
    ]
    bullets(slide, Inches(1.0), Inches(5.05), Inches(5.2), Inches(0.8),
            mode1, size=11, color=LIGHT_GRAY, spacing=1.5)

    rect(slide, Inches(6.9), Inches(4.5), Inches(5.8), Inches(1.5), DARK_CARD, radius=0.06)
    textbox(slide, Inches(7.2), Inches(4.65), Inches(3.0), Inches(0.35),
            "🔸 模式二：任意文件分析", size=15, color=ACCENT_ORANGE, bold=True)
    mode2 = [
        "• .xlsx / .xls / .csv → 智能列检测 → 报告/图表",
        "• 适用：临时数据探索，无需配置，拖入即分析",
    ]
    bullets(slide, Inches(7.2), Inches(5.05), Inches(5.2), Inches(0.8),
            mode2, size=11, color=LIGHT_GRAY, spacing=1.5)

    # Tech stack tags
    textbox(slide, Inches(0.8), Inches(6.4), Inches(1.5), Inches(0.3),
            "技术栈：", size=12, color=LIGHT_GRAY)
    tags_data = [
        ("Python 3.11+", ACCENT_BLUE), ("pandas", ACCENT_TEAL),
        ("pymssql", ACCENT_GREEN), ("Click CLI", ACCENT_PURPLE),
        ("openpyxl", ACCENT_ORANGE), ("matplotlib", ACCENT_BLUE),
        ("reportlab", ACCENT_TEAL), ("structlog", ACCENT_GREEN),
    ]
    for i, (t, c) in enumerate(tags_data):
        tag(slide, Inches(2.5) + i * Inches(1.3), Inches(6.35), t, c)

    slide_num(slide, 3)


# ═══════════════════════════════════════════════
# SLIDE 4 — Feature: Config-driven Export
# ═══════════════════════════════════════════════
def slide_feature_export(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, Inches(0), Inches(0), Inches(13.33), Inches(7.5), BG_DARK)
    header_bar(slide, "功能一：配置化数据导出", "SQL Server → .xlsx  —  YAML 驱动，零代码，安全可靠")

    cards_data = [
        ("📝 YAML 配置驱动", ACCENT_BLUE, [
            "• 服务器/端口/数据库/SQL 全部 YAML 化",
            "• 支持多查询组，每组独立覆盖连接参数",
            "• detail_query + summary_query 双查询",
            "• 人工可读写，支持注释",
        ]),
        ("🔐 密码安全机制", ACCENT_GREEN, [
            "• 强制 ${ENV_VAR} 环境变量引用",
            "• 拒绝明文密码 — 配置可安全提交 Git",
            "• 配置 Schema 严格校验（端口1-65535）",
            "• 错误配置 5 秒内返回明确提示",
        ]),
        ("🔄 参数化查询", ACCENT_PURPLE, [
            "• SQL 中 {{变量}} 占位符自动替换",
            "• 支持：日期/企业ID/商品ID/员工ID",
            "• Prepared Statement 防 SQL 注入",
            "• 日期范围以天为最小粒度 (YYYY-MM-DD)",
        ]),
        ("📤 自动导出 .xlsx", ACCENT_ORANGE, [
            "• pandas + openpyxl 引擎，UTF-8 编码",
            "• 自动列宽调整，时间戳命名",
            "• 查询结果为空 → 记录日志，不报错",
            "• 连接失败 → 明确错误信息，不产生脏数据",
        ]),
    ]

    cw = Inches(5.8)
    ch = Inches(2.35)
    for i, (title, color, items) in enumerate(cards_data):
        col, row = i % 2, i // 2
        x = Inches(0.55) + col * (cw + Inches(0.35))
        y = Inches(1.6) + row * (ch + Inches(0.2))
        card(slide, x, y, cw, ch, title, items, title_color=color)

    # Example command
    rect(slide, Inches(0.55), Inches(6.65), Inches(12.2), Inches(0.55), DARK_CARD, radius=0.04)
    textbox(slide, Inches(0.8), Inches(6.72), Inches(11.7), Inches(0.4),
            "📌 命令:  python -m src.cli export   |   python -m src.cli export -q sanzhen-jiuti   |   python -m src.cli run（一键全流程）",
            size=12, color=ACCENT_BLUE, font="Consolas")
    slide_num(slide, 4)


# ═══════════════════════════════════════════════
# SLIDE 5 — Feature: Smart Template Analysis
# ═══════════════════════════════════════════════
def slide_feature_analysis(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, Inches(0), Inches(0), Inches(13.33), Inches(7.5), BG_DARK)
    header_bar(slide, "功能二：智能模版分析引擎", "自动匹配 · 分组汇总 · 占比计算 — 100% 计算精确")

    # Left — Auto template generation
    rect(slide, Inches(0.5), Inches(1.6), Inches(6.0), Inches(2.6), DARK_CARD, radius=0.06)
    textbox(slide, Inches(0.8), Inches(1.75), Inches(5.4), Inches(0.4),
            "🧠  自动模版生成（零手工配置）", size=18, color=ACCENT_BLUE, bold=True)
    auto_items = [
        "• 从 SQL 查询结果自动检测所有列结构",
        "• 智能识别：金额列（含「金额」字样）→ 名称列（含「姓名」字样）",
        "• 根据 group_by 配置提取 片区/门店/员工/部门 信息",
        "• match_key 指定员工 ID 匹配字段，自动去重生成员工列表",
        "• 无需手写 78 行员工 YAML — 零维护成本",
    ]
    bullets(slide, Inches(0.8), Inches(2.3), Inches(5.4), Inches(1.7),
            auto_items, size=11.5, color=LIGHT_GRAY, spacing=1.55)

    # Right — Analysis logic
    rect(slide, Inches(6.85), Inches(1.6), Inches(6.0), Inches(2.6), DARK_CARD, radius=0.06)
    textbox(slide, Inches(7.15), Inches(1.75), Inches(5.4), Inches(0.4),
            "📐  分析计算逻辑", size=18, color=ACCENT_GREEN, bold=True)
    calc_items = [
        "• 按 门店→员工 维度 GROUP BY 汇总销售金额",
        "• 占比 = 员工销售额 ÷ 连锁月度总销售额 × 100%",
        "• 合计行与 SQL 直接查询总额偏差 = 0（精确匹配）",
        "• 模版有但数据无 → 销售额显示为 ¥0.00",
        "• 数据有但模版无 → 附录单独列出，不计入合计",
    ]
    bullets(slide, Inches(7.15), Inches(2.3), Inches(5.4), Inches(1.7),
            calc_items, size=11.5, color=LIGHT_GRAY, spacing=1.55)

    # Bottom — Analysis flow
    textbox(slide, Inches(0.6), Inches(4.5), Inches(3.0), Inches(0.35),
            "🔄  分析处理流程", size=16, color=WHITE, bold=True)
    flow = [
        ("加载\nDataFrame", ACCENT_BLUE), ("自动生成\n模版", ACCENT_PURPLE),
        ("员工匹配\n& 汇总", ACCENT_TEAL), ("占比\n计算", ACCENT_ORANGE),
        ("输出\nAnalysisResult", ACCENT_GREEN),
    ]
    fw = Inches(2.1)
    fx = Inches(0.8)
    for i, (label, color) in enumerate(flow):
        x = fx + i * (fw + Inches(0.3))
        rect(slide, x, Inches(5.0), fw, Inches(1.0), color, radius=0.08)
        textbox(slide, x, Inches(5.1), fw, Inches(0.8),
                label, size=13, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        if i < len(flow) - 1:
            textbox(slide, x + fw, Inches(5.15), Inches(0.3), Inches(0.5),
                    "→", size=22, color=LIGHT_GRAY, align=PP_ALIGN.CENTER)

    # Metrics
    rect(slide, Inches(0.5), Inches(6.45), Inches(12.3), Inches(0.7), DARK_CARD, radius=0.04)
    metrics_text = "匹配覆盖率 100%  |  合计偏差 0 元  |  支持 10 万行级数据  |  每员工精确验证通过  |  78 人模版自动生成"
    textbox(slide, Inches(0.8), Inches(6.55), Inches(11.7), Inches(0.5),
            metrics_text, size=13, color=ACCENT_BLUE, align=PP_ALIGN.CENTER)
    slide_num(slide, 5)


# ═══════════════════════════════════════════════
# SLIDE 6 — Feature: Multi-format Reports
# ═══════════════════════════════════════════════
def slide_feature_report(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, Inches(0), Inches(0), Inches(13.33), Inches(7.5), BG_DARK)
    header_bar(slide, "功能三：多格式报告生成", "Markdown + Excel + PDF — 一次分析，三种交付格式")

    formats = [
        ("📝  Markdown 报告", ACCENT_BLUE, [
            "• 标题含日期范围，结构化排版",
            "• 汇总统计（总销售额/人数/门店数/匹配率）",
            "• 详细分析表格（与模版格式完全一致）",
            "• 合计行 + 未匹配数据附录",
            "• UTF-8 编码，可直接阅读/编辑/归档",
        ]),
        ("📊  Excel 报告", ACCENT_GREEN, [
            "• openpyxl 格式化 — 表头加粗+蓝底",
            "• 金额列 #,##0.00 千分位格式",
            "• 百分比列 0.00% 格式自动转换",
            "• 0 销售额员工灰底标注区分",
            "• 自动列宽 + 细线边框 + 合并标题行",
        ]),
        ("📄  PDF 报告", ACCENT_ORANGE, [
            "• reportlab 生成，A4 横向排版",
            "• 微软雅黑中文字体完美嵌入",
            "• 标题页 + 汇总统计 + 数据表格 + 图表",
            "• 柱状图/饼图/漏斗图嵌入正文",
            "• 适合打印、客户汇报、正式归档",
        ]),
    ]

    cw = Inches(3.85)
    for i, (title, color, items) in enumerate(formats):
        x = Inches(0.55) + i * (cw + Inches(0.25))
        card(slide, x, Inches(1.6), cw, Inches(3.15), title, items, title_color=color)

    # Output structure
    rect(slide, Inches(0.55), Inches(5.0), Inches(12.2), Inches(2.2), DARK_CARD, radius=0.06)
    textbox(slide, Inches(0.8), Inches(5.1), Inches(4.0), Inches(0.35),
            "📁  输出目录结构", size=15, color=WHITE, bold=True)
    structure = [
        "output/reports/",
        "  ├── export_sanzhen-jiuti_20260626/",
        "  │   ├── chart_bar_*.png          ← 柱状图",
        "  │   ├── chart_pie_*.png          ← 饼图",
        "  │   ├── chart_funnel_*.png       ← 漏斗图",
        "  │   ├── report.xlsx              ← Excel 报告",
        "  │   └── report.pdf               ← PDF 报告（可选）",
    ]
    bullets(slide, Inches(0.8), Inches(5.5), Inches(5.5), Inches(1.5),
            structure, size=10.5, color=LIGHT_GRAY, spacing=1.3)

    textbox(slide, Inches(7.0), Inches(5.1), Inches(5.5), Inches(0.35),
            "🎯  报告特点", size=15, color=WHITE, bold=True)
    features = [
        "• 全自动生成，无需任何手工排版",
        "• 所有文件可直接打开使用",
        "• 图表 300 DPI 高清 PNG，适合放大查看",
        "• 支持 --no-pdf / --no-charts 按需生成",
    ]
    bullets(slide, Inches(7.0), Inches(5.5), Inches(5.5), Inches(1.5),
            features, size=11, color=LIGHT_GRAY, spacing=1.65)

    slide_num(slide, 6)


# ═══════════════════════════════════════════════
# SLIDE 7 — Feature: Charts
# ═══════════════════════════════════════════════
def slide_feature_charts(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, Inches(0), Inches(0), Inches(13.33), Inches(7.5), BG_DARK)
    header_bar(slide, "功能四：可视化图表引擎", "Matplotlib 渲染 · 中文字体自动发现 · 300 DPI 高清输出")

    charts_data = [
        ("📊  柱状图 — Top 10 销售排行", ACCENT_BLUE, [
            "• 水平条形图，自动取 Top 10 + 其他",
            "• 数值标签显示 ¥ 金额，一目了然",
            "• 适合：员工业绩排行、门店对比",
        ]),
        ("🥧  饼图 — 销售结构分布", ACCENT_GREEN, [
            "• 按片区/门店自动聚合分组",
            "• < 3% 小类归入「其他」类别",
            "• Set3 专业配色 + 侧边图例",
            "• 适合：区域占比、品类结构分析",
        ]),
        ("🔽  漏斗图 — Top 10 业绩漏斗", ACCENT_ORANGE, [
            "• 渐变色漏斗，最大金额在最顶部",
            "• 左侧标签 + 右侧金额，紧凑布局",
            "• 无坐标轴极简设计，视觉聚焦",
            "• 适合：排名分布、业绩梯度展示",
        ]),
    ]

    for i, (title, color, items) in enumerate(charts_data):
        y = Inches(1.7) + i * Inches(1.6)
        rect(slide, Inches(0.5), y, Inches(0.08), Inches(1.3), color)
        textbox(slide, Inches(0.9), y, Inches(7.0), Inches(0.4),
                title, size=18, color=color, bold=True)
        bullets(slide, Inches(0.9), y + Inches(0.45), Inches(6.5), Inches(0.8),
                items, size=12, color=LIGHT_GRAY, spacing=1.5)

    # Right — tech details
    rect(slide, Inches(8.0), Inches(1.7), Inches(4.8), Inches(5.0), DARK_CARD, radius=0.06)
    textbox(slide, Inches(8.3), Inches(1.85), Inches(4.2), Inches(0.4),
            "🛠  技术细节", size=16, color=ACCENT_PURPLE, bold=True)
    tech = [
        "• Matplotlib Agg 后端（无 GUI）",
        "• 自动发现系统中文字体：",
        "    Windows: 微软雅黑 / 黑体 / 宋体",
        "    Linux: Noto / WenQuanYi",
        "• 300 DPI PNG 高清渲染",
        "• 图表可嵌入 PDF 报告正文",
        "• 纯后端 — 无需浏览器或 JS",
        "• generate_charts_for_analysis()",
        "   — 一行调用，3 张图秒出",
        "• generate_charts_for_file()",
        "   — 任意文件即可生成图表",
    ]
    bullets(slide, Inches(8.3), Inches(2.35), Inches(4.2), Inches(3.5),
            tech, size=10.5, color=LIGHT_GRAY, spacing=1.5)

    slide_num(slide, 7)


# ═══════════════════════════════════════════════
# SLIDE 8 — Feature: File Analysis
# ═══════════════════════════════════════════════
def slide_feature_file(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, Inches(0), Inches(0), Inches(13.33), Inches(7.5), BG_DARK)
    header_bar(slide, "功能五：任意文件智能分析", "analyze-file — 无需 SQL · 无需模版 · 拖入即分析")

    # Left — How it works
    rect(slide, Inches(0.5), Inches(1.6), Inches(6.0), Inches(3.8), DARK_CARD, radius=0.06)
    textbox(slide, Inches(0.8), Inches(1.75), Inches(5.4), Inches(0.35),
            "🔍  智能检测流程", size=17, color=ACCENT_BLUE, bold=True)
    detect = [
        "1. 加载文件（自动编码检测 UTF-8→GBK→Latin-1）",
        "2. 检测数值列 — 优先含「金额」列，否则取最大和数值列",
        "3. 检测名称列 — 优先「姓名」→「名称」→ 首字符串列",
        "4. 检测分组列 — 优先「门店」→「片区」→ 最优分类列",
        "5. 按数值降序排列，自动计算每行占比",
        "6. 生成 Excel 报告 + 图表 PNG（可选 PDF）",
    ]
    bullets(slide, Inches(0.8), Inches(2.2), Inches(5.4), Inches(3.0),
            detect, size=12, color=LIGHT_GRAY, spacing=1.6)

    # Right — Supported formats & manual override
    rect(slide, Inches(6.85), Inches(1.6), Inches(6.0), Inches(1.6), DARK_CARD, radius=0.06)
    textbox(slide, Inches(7.15), Inches(1.75), Inches(5.4), Inches(0.35),
            "📂  支持格式", size=17, color=ACCENT_GREEN, bold=True)
    formats = [
        "• .xlsx  — Excel 2007+ 格式",
        "• .xls   — Excel 97-2003 格式",
        "• .csv   — 逗号分隔，自动检测编码",
        "• 单文件或整个目录批量处理",
    ]
    bullets(slide, Inches(7.15), Inches(2.2), Inches(5.4), Inches(0.9),
            formats, size=12, color=LIGHT_GRAY, spacing=1.4)

    rect(slide, Inches(6.85), Inches(3.45), Inches(6.0), Inches(1.95), DARK_CARD, radius=0.06)
    textbox(slide, Inches(7.15), Inches(3.6), Inches(5.4), Inches(0.35),
            "⚙️  高级选项（可选手动指定）", size=17, color=ACCENT_ORANGE, bold=True)
    opts = [
        "• --value-col \"金额\"    — 手动指定数值列",
        "• --name-col \"姓名\"     — 手动指定名称列",
        "• --group-col \"区域\"    — 手动指定分组列",
        "• --top-n 15             — 图表 Top N 数量",
        "• --no-pdf --no-charts   — 仅生成 Excel",
    ]
    bullets(slide, Inches(7.15), Inches(4.05), Inches(5.4), Inches(1.2),
            opts, size=11, color=LIGHT_GRAY, spacing=1.4)

    # Bottom — use cases
    rect(slide, Inches(0.5), Inches(5.65), Inches(12.3), Inches(1.55), DARK_CARD, radius=0.06)
    textbox(slide, Inches(0.8), Inches(5.75), Inches(4.0), Inches(0.35),
            "🎯  适用场景", size=16, color=WHITE, bold=True)
    use_cases = [
        "• 临时数据探索 — 拿到一份 Excel 就想看排名/分布/图表                         • 批量处理 — 导出目录下所有文件一键生成全套报告",
        "• 无需配置 — 不写 YAML、不写 SQL，拖入即分析                                  • 快速验证 — 检查任意数据文件的结构和质量",
    ]
    bullets(slide, Inches(0.8), Inches(6.15), Inches(11.7), Inches(0.9),
            use_cases, size=11.5, color=LIGHT_GRAY, spacing=1.5)
    slide_num(slide, 8)


# ═══════════════════════════════════════════════
# SLIDE 9 — Technical Architecture
# ═══════════════════════════════════════════════
def slide_tech(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, Inches(0), Inches(0), Inches(13.33), Inches(7.5), BG_DARK)
    header_bar(slide, "技术架构与模块设计", "模块化 · 单一职责 · 易扩展 · 全测试覆盖")

    # Module structure
    rect(slide, Inches(0.5), Inches(1.6), Inches(6.0), Inches(3.0), DARK_CARD, radius=0.06)
    textbox(slide, Inches(0.8), Inches(1.7), Inches(5.4), Inches(0.35),
            "📦  源码模块结构 (src/)", size=16, color=ACCENT_BLUE, bold=True)
    modules = [
        "src/config/      配置加载 & Schema 校验",
        "src/export/      SQL Server 连接 & 数据导出",
        "src/analysis/    模版匹配 & 分组汇总计算",
        "src/report/      Markdown / Excel / PDF 生成",
        "src/template/    模版注册 & 自动生成引擎",
        "src/chart/       柱状图 / 饼图 / 漏斗图渲染",
        "src/pipeline.py          全流程编排引擎",
        "src/pipeline_file.py     任意文件分析管道",
        "src/cli.py               Click CLI 统一入口",
    ]
    bullets(slide, Inches(0.8), Inches(2.15), Inches(5.4), Inches(2.3),
            modules, size=11, color=LIGHT_GRAY, spacing=1.4, font_name="Consolas")

    # Tech stack
    rect(slide, Inches(6.85), Inches(1.6), Inches(6.0), Inches(3.0), DARK_CARD, radius=0.06)
    textbox(slide, Inches(7.15), Inches(1.7), Inches(5.4), Inches(0.35),
            "🛠  核心技术栈", size=16, color=ACCENT_GREEN, bold=True)
    stack = [
        "Python 3.11+        — 主语言",
        "Click                — CLI 命令框架",
        "pandas               — 数据计算引擎",
        "pymssql              — SQL Server 连接",
        "openpyxl             — Excel 读写生成",
        "reportlab            — PDF 文档生成",
        "matplotlib           — 图表渲染（Agg）",
        "structlog            — 结构化日志",
        "pytest (41 cases)    — TDD 测试",
    ]
    bullets(slide, Inches(7.15), Inches(2.15), Inches(5.4), Inches(2.3),
            stack, size=11, color=LIGHT_GRAY, spacing=1.4)

    # Constitution & Tests
    rect(slide, Inches(0.5), Inches(4.85), Inches(12.3), Inches(2.35), DARK_CARD, radius=0.06)
    textbox(slide, Inches(0.8), Inches(4.95), Inches(5.5), Inches(0.35),
            "✅  项目宪法合规", size=16, color=ACCENT_GREEN, bold=True)
    const = [
        "I. 规格优先 — spec.md / plan.md / tasks.md 完整体系     IV. 精准修改 — 纯新增功能，无现有代码破坏",
        "II. 简单至上 — 每模块职责单一，无过度设计                V. 目标驱动 — TDD 测试先行，Red-Green-Refactor",
        "III. 先想后写 — 所有技术决策显式记录                     VI. TDD 铁律 — 41 个用例全部通过",
    ]
    bullets(slide, Inches(0.8), Inches(5.35), Inches(11.7), Inches(1.1),
            const, size=10.5, color=LIGHT_GRAY, spacing=1.5)

    textbox(slide, Inches(0.8), Inches(6.4), Inches(11.7), Inches(0.4),
            "🧪  测试: 8 个测试模块 × 41 个用例  |  边界场景全覆盖: 5000 行性能 / NULL 值 / 类型归一化 / 日期边界 / 密码安全 / SQL 注入防护",
            size=11.5, color=ACCENT_BLUE)
    slide_num(slide, 9)


# ═══════════════════════════════════════════════
# SLIDE 10 — CLI Commands
# ═══════════════════════════════════════════════
def slide_cli(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, Inches(0), Inches(0), Inches(13.33), Inches(7.5), BG_DARK)
    header_bar(slide, "CLI 命令一览", "统一命令行入口 — 5 个核心命令覆盖全部功能")

    commands = [
        ("run", "一键全流程", ACCENT_BLUE,
         "python -m src.cli run [--template] [--auto-date] [--no-pdf] [--no-charts]",
         "导出 → 分析 → 报告 → 图表，全自动完成。支持 --auto-date 定时调度（每月1号/16号）"),
        ("export", "仅数据导出", ACCENT_TEAL,
         "python -m src.cli export [-c config.yaml] [-q 查询组]",
         "执行 SQL 查询 → .xlsx 文件，支持单独导出指定查询组"),
        ("analyze-file", "任意文件分析", ACCENT_ORANGE,
         "python -m src.cli analyze-file <文件|目录> [--value-col] [--group-col] [--top-n]",
         "智能分析 .xlsx/.xls/.csv，自动检测列结构生成报告+图表"),
        ("show-queries", "查看查询组", ACCENT_GREEN,
         "python -m src.cli show-queries [-c config.yaml]",
         "列出 export.yaml 中所有查询组及其连接/模版/参数配置摘要"),
        ("validate", "配置校验", ACCENT_PURPLE,
         "python -m src.cli validate [-c config.yaml]",
         "校验配置文件格式完整性、连接参数合法性、SQL 安全性"),
    ]

    for i, (cmd, title, color, example, desc) in enumerate(commands):
        y = Inches(1.7) + i * Inches(1.05)
        rect(slide, Inches(0.5), y, Inches(12.3), Inches(0.85), DARK_CARD, radius=0.04)
        rect(slide, Inches(0.5), y, Inches(0.06), Inches(0.85), color)
        # Command name badge
        rect(slide, Inches(0.8), y + Inches(0.15), Inches(1.5), Inches(0.28), color, radius=0.12)
        textbox(slide, Inches(0.8), y + Inches(0.15), Inches(1.5), Inches(0.28),
                cmd, size=11, color=WHITE, bold=True, align=PP_ALIGN.CENTER, font="Consolas")
        textbox(slide, Inches(2.5), y + Inches(0.12), Inches(2.0), Inches(0.3),
                title, size=14, color=color, bold=True)
        textbox(slide, Inches(4.5), y + Inches(0.12), Inches(7.5), Inches(0.25),
                example, size=9.5, color=ACCENT_BLUE, font="Consolas")
        textbox(slide, Inches(0.8), y + Inches(0.48), Inches(11.5), Inches(0.3),
                desc, size=10.5, color=LIGHT_GRAY)

    slide_num(slide, 10)


# ═══════════════════════════════════════════════
# SLIDE 11 — Project Highlights
# ═══════════════════════════════════════════════
def slide_highlights(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, Inches(0), Inches(0), Inches(13.33), Inches(7.5), BG_DARK)
    header_bar(slide, "项目亮点", "安全性 · 零人工 · 可扩展 · 生产可靠")

    highlights = [
        ("🔐", "安全第一", ACCENT_BLUE, [
            "密码强制 ${ENV_VAR} 环境变量引用",
            "配置文件可安全提交版本控制",
            "Prepared Statement 防 SQL 注入",
            "Schema 校验拒绝非法配置",
        ]),
        ("⚡", "零人工介入", ACCENT_GREEN, [
            "SQL 导出 → 报告生成全自动",
            "模版从查询结果自动生成",
            "analyze-file 拖入即分析",
            "定时调度 --auto-date 无人值守",
        ]),
        ("🔌", "高度可扩展", ACCENT_PURPLE, [
            "多查询组并行管理架构",
            "每组独立连接/SQL/模版参数",
            "热加载模版注册表",
            "新增业务场景只需加配置",
        ]),
        ("🛡", "生产级可靠", ACCENT_ORANGE, [
            "41 个 TDD 测试全部通过",
            "5000 行 < 30 秒性能基准",
            "structlog 全链路结构化日志",
            "异常翻译为中文可读错误",
        ]),
    ]

    cw = Inches(5.9)
    ch = Inches(2.3)
    for i, (icon, title, color, items) in enumerate(highlights):
        col, row = i % 2, i // 2
        x = Inches(0.55) + col * (cw + Inches(0.35))
        y = Inches(1.6) + row * (ch + Inches(0.25))
        rect(slide, x, y, cw, ch, DARK_CARD, radius=0.06)
        rect(slide, x, y, cw, Inches(0.05), color)
        textbox(slide, x + Inches(0.2), y + Inches(0.15), cw - Inches(0.4), Inches(0.4),
                f"{icon}  {title}", size=17, color=color, bold=True)
        bullets(slide, x + Inches(0.2), y + Inches(0.7), cw - Inches(0.4), ch - Inches(0.9),
                items, size=11.5, color=LIGHT_GRAY, spacing=1.6)

    # Bottom stat bar
    rect(slide, Inches(0.5), Inches(6.55), Inches(12.3), Inches(0.65), DARK_CARD, radius=0.04)
    stats = "41 个测试  |  100% 覆盖率  |  8 个模块  |  < 30 秒全流程  |  0 计算偏差  |  10+ 模版容量"
    textbox(slide, Inches(0.8), Inches(6.62), Inches(11.7), Inches(0.5),
            stats, size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    slide_num(slide, 11)


# ═══════════════════════════════════════════════
# SLIDE 12 — Summary & Thanks
# ═══════════════════════════════════════════════
def slide_summary(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, Inches(0), Inches(0), Inches(13.33), Inches(7.5), BG_DARK)
    header_bar(slide, "总结与展望", "AIExport — 让数据导出分析像呼吸一样自然")

    # Left — Deliverables
    rect(slide, Inches(0.5), Inches(1.6), Inches(6.0), Inches(3.6), DARK_CARD, radius=0.06)
    textbox(slide, Inches(0.8), Inches(1.7), Inches(5.4), Inches(0.35),
            "📋  项目交付成果", size=18, color=ACCENT_GREEN, bold=True)
    deliverables = [
        "✅  配置化 SQL Server 数据导出",
        "✅  智能模版自动生成与匹配分析",
        "✅  Markdown + Excel + PDF 三格式报告",
        "✅  柱状图/饼图/漏斗图可视化",
        "✅  任意文件智能分析（xlsx/xls/csv）",
        "✅  多查询组 + 定时调度 + 自动日期",
        "✅  41 个 TDD 测试用例全部通过",
        "✅  完整 spec / plan / tasks 文档体系",
    ]
    bullets(slide, Inches(0.8), Inches(2.15), Inches(5.4), Inches(2.8),
            deliverables, size=12, color=LIGHT_GRAY, spacing=1.6)

    # Right — Roadmap
    rect(slide, Inches(6.85), Inches(1.6), Inches(6.0), Inches(3.6), DARK_CARD, radius=0.06)
    textbox(slide, Inches(7.15), Inches(1.7), Inches(5.4), Inches(0.35),
            "🚀  后续规划", size=18, color=ACCENT_ORANGE, bold=True)
    roadmap = [
        "🔲  HTML 报告格式 — Web 浏览器查看",
        "🔲  Web Dashboard — 实时监控面板",
        "🔲  邮件自动发送 — 定时推送到邮箱",
        "🔲  更多数据源 — MySQL / PostgreSQL",
        "🔲  AI 智能洞察 — 异常检测/趋势预测",
        "🔲  Webhook 集成 — 企业微信/钉钉通知",
        "🔲  数据对比 — 同比/环比自动分析",
        "🔲  Cron 定时 — 操作系统级任务调度",
    ]
    bullets(slide, Inches(7.15), Inches(2.15), Inches(5.4), Inches(2.8),
            roadmap, size=11.5, color=LIGHT_GRAY, spacing=1.5)

    # Bottom — value proposition
    rect(slide, Inches(0.5), Inches(5.5), Inches(12.3), Inches(1.7), ACCENT_BLUE, radius=0.06)
    textbox(slide, Inches(0.8), Inches(5.65), Inches(11.7), Inches(0.35),
            "💎  核心价值", size=14, color=RGBColor(0xBB, 0xCC, 0xEE))
    textbox(slide, Inches(0.8), Inches(5.95), Inches(11.7), Inches(0.55),
            "将 2~4 小时的手工流程压缩至 30 秒内全自动完成",
            size=24, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    textbox(slide, Inches(0.8), Inches(6.45), Inches(11.7), Inches(0.4),
            "零人工介入 · 100% 计算精确 · 多格式报告 · 可视化图表 · 安全可靠",
            size=14, color=RGBColor(0xDD, 0xDD, 0xFF), align=PP_ALIGN.CENTER)

    # Thank you
    textbox(slide, Inches(0.8), Inches(7.05), Inches(11.7), Inches(0.3),
            "感谢聆听  ·  欢迎提问", size=14, color=LIGHT_GRAY, align=PP_ALIGN.CENTER)
    slide_num(slide, 12)


# ═══════════════════════════════════════════════
# Main
# ═══════════════════════════════════════════════
def main():
    prs = Presentation()
    prs.slide_width  = Inches(13.33)  # 16:9 widescreen
    prs.slide_height = Inches(7.5)

    slide_cover(prs)
    slide_background(prs)
    slide_architecture(prs)
    slide_feature_export(prs)
    slide_feature_analysis(prs)
    slide_feature_report(prs)
    slide_feature_charts(prs)
    slide_feature_file(prs)
    slide_tech(prs)
    slide_cli(prs)
    slide_highlights(prs)
    slide_summary(prs)

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    prs.save(PPT_PATH)
    print(f"[OK] PPT saved: {PPT_PATH}")
    print(f"     {len(prs.slides)} slides, 16:9 format")


if __name__ == "__main__":
    main()
